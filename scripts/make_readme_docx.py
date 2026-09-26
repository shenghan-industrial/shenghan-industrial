# -*- coding: utf-8 -*-
"""Generate a Word instruction sheet for the DEXOREN site source package."""
from docx import Document
from docx.shared import Pt, RGBColor, Cm

OUT = "/Users/yiyi/Desktop/DEXOREN独立站-使用说明.docx"

doc = Document()
st = doc.styles["Normal"]
st.font.name = "Arial Unicode MS"
st.font.size = Pt(10.5)


def h(text, level=1):
    p = doc.add_heading(text, level=level)
    for r in p.runs:
        r.font.color.rgb = RGBColor(0x3D, 0x37, 0x30)
    return p


def para(text, bold=False, italic=False):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = bold
    r.italic = italic
    return p


def code(text):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.font.name = "Courier New"
    r.font.size = Pt(9.5)
    r.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
    p.paragraph_format.left_indent = Cm(0.6)
    p.paragraph_format.space_after = Pt(4)
    return p


def bullet(text):
    return doc.add_paragraph(text, style="List Bullet")


t = doc.add_heading("DEXOREN 独立站源码包", level=0)
for r in t.runs:
    r.font.color.rgb = RGBColor(0x3D, 0x37, 0x30)
para("在另一台电脑上运行本站点 —— 使用说明", italic=True)
doc.add_paragraph()

h("一、这个包里有什么", 1)
bullet("完整网站源码（Next.js 15 项目），品牌名为 DEXOREN")
bullet("产品数据：2,343 款产品（data/products.json 与运行时 .data/products.json）")
bullet("产品图片：public/images、public/uploads 全部图片资源")
bullet("6 种语言文案：中文 / English / ไทย(泰语) / Bahasa Melayu(马来语) / Français(法语) / Español(西班牙语)")
bullet("产品管理后台（/admin）与上架脚本（scripts/）")
bullet("另有单独的产品目录 PDF：DEXOREN-Product-Catalog.pdf（可一并拷贝）")

h("不包含（这两项可再生，省了约 900 MB）", 2)
bullet("node_modules —— 依赖目录，解压后执行 npm install 自动还原")
bullet(".next —— 构建缓存，启动后自动生成")

h("二、另一台电脑需要准备什么", 1)
bullet("Node.js 18 或更高版本（推荐 20 或 22）")
code("官网下载：https://nodejs.org  →  选 LTS 版本安装即可")
bullet("首次安装依赖需要联网")
bullet("系统：macOS / Windows / Linux 都可以")

h("三、三步启动（在另一台电脑上）", 1)
para("第 1 步：解压", bold=True)
code("把 DEXOREN-site-source.zip 解压到任意位置")
para("解压后会得到一个 dexoren-site 文件夹。", italic=True)

para("第 2 步：安装依赖", bold=True)
para("打开终端（Windows 用 PowerShell），进入解压后的文件夹：")
code("cd dexoren-site")
code("npm install")
para("大约 1–3 分钟，看到没有红色报错即成功。", italic=True)

para("第 3 步：启动网站", bold=True)
code("npm run dev")
para("启动后浏览器打开：")
code("http://localhost:3000")
para("（若 3000 端口被占用，会自动换成 3001、3002……看终端提示的地址）")

h("四、常用地址", 1)
table = doc.add_table(rows=1, cols=2)
table.style = "Light Grid Accent 1"
hdr = table.rows[0].cells
hdr[0].text = "页面"
hdr[1].text = "地址"
for a, b in [
    ("网站首页", "http://localhost:3000/"),
    ("产品列表", "http://localhost:3000/products"),
    ("贸易条款", "http://localhost:3000/trade-terms"),
    ("认证资质", "http://localhost:3000/certifications"),
    ("关于我们", "http://localhost:3000/about"),
    ("联系我们", "http://localhost:3000/contact"),
    ("管理后台", "http://localhost:3000/admin"),
]:
    c = table.add_row().cells
    c[0].text = a
    c[1].text = b

doc.add_paragraph()

h("五、重要提醒", 1)

para("1. 关于品牌与域名", bold=True)
bullet("品牌名已统一为 DEXOREN")
bullet("网站域名与联系邮箱暂时沿用 shenghanindustrial.com（sales@shenghanindustrial.com）")
bullet("若后续启用新域名，需同时改：邮箱、页脚、SEO 结构化数据（lib/schema-org.tsx）")

para("2. 关于产品图片", bold=True)
bullet("图片分两处：本地 public/ 目录 + Cloudinary 云图床")
bullet("Cloudinary 上的图是公网直链，换电脑不影响；本地图片已全部打包")
bullet("Cloudinary 凭据写在 lib/cloudinary.ts，换图床改这里")

para("3. 关于产品数据", bold=True)
bullet("数据存两份：data/products.json（源）与 .data/products.json（运行时）")
bullet("后台增删改产品写入 .data/products.json")
bullet("要长期保存或部署上线，需把 .data 内容同步回 data/products.json")

para("4. 关于语言", bold=True)
bullet("右上角可切换 6 种语言；新增语言只需在 messages/ 加对应 json 文件")

h("六、常见问题", 1)
para("npm install 报错或很慢？", bold=True)
bullet("可换成国内镜像后重装：")
code("npm config set registry https://registry.npmmirror.com")

para("提示端口被占用？", bold=True)
code("npm run dev -- -p 3002")

para("想部署成正式网站？", bold=True)
code("npm run build")
bullet("构建通过后部署到服务器或 Cloudflare Pages")

para("npm 命令不存在？", bold=True)
bullet("说明 Node.js 没装好，重新安装后重开终端")

doc.add_paragraph()
para("—— 有任何问题随时联系即可 ——", italic=True)

doc.save(OUT)
print("saved:", OUT)
